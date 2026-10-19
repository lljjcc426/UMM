"""Regenerate Markdown, CSV and XLSX views from the canonical literature catalog."""
from pathlib import Path
import argparse
import csv
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
HEADER = ["编号","类别","研究论文","版本","声明","方法","针对问题","结论","缺口","原始来源","核查状态","最近原文核查日","历史核查说明","历史编号",
          "数据集与分布","方法框架","模型结构","实验指标","对比方法","实验条件","阅读笔记路径","本批阅读层级"]
VERIFICATION_LABELS = {
    "historical_record_not_rechecked": "历史记录，待复核",
    "primary_sections_read": "相关原文章节已核查",
    "primary_abstract_read": "原始摘要已核查",
    "primary_metadata_only": "仅核查元数据",
    "historical_record_unresolved": "原始来源待核",
}

def cells(r):
    vals = [r["id"],r["category"],r["title"],r["version"],r["claim"],r["method"],r["problem"],
            r["conclusion"],r["gap"],"\n".join(r["source_urls"]),r["verification"],
            r.get("checked_on") or "",r.get("previous_evidence") or "",
            "\n".join(r.get("aliases",[]))] + [r.get(k, "") for k in
            ("dataset", "framework", "architecture", "metrics", "baselines", "experimental_conditions")]
    vals += [r.get("review_path", r.get("note_path", "")), r.get("review_depth", "")]
    return ["'"+v if isinstance(v,str) and v.startswith(("=","+","-","@")) else v for v in vals]

def md(recs, title, destination):
    lines = [f"# {title}", "", "按六字段组织；作者报告不等于独立复现。每行标明核查状态。",
             "", "| 研究论文 | 声明 | 方法 | 针对问题 | 结论 | 缺口 |",
             "| --- | --- | --- | --- | --- | --- |"]
    for r in recs:
        lead = f'**{r["id"]} · {r["title"]}**<br>{r["version"]}<br>'
        lead += " / ".join(f"[来源{i+1}]({u})" for i,u in enumerate(r["source_urls"]))
        lead += f'<br>核查：{VERIFICATION_LABELS[r["verification"]]}'
        if r.get("review_path"):
            note = os.path.relpath(ROOT / r["review_path"], (ROOT / destination).parent)
            lead += f'<br>[结构化阅读与实验条件]({note})'
        vals = [lead] + [r[k] for k in ("claim","method","problem","conclusion","gap")]
        lines.append("| "+" | ".join(str(v).replace("|",r"\|").replace("\n","<br>") for v in vals)+" |")
    return "\n".join(lines)+"\n"

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-xlsx", action="store_true", help="Render text views only; preserve the current workbook for a separate editor.")
    parser.add_argument("--xlsx-payload", type=Path, help="Write worksheet matrices for a separate XLSX editor.")
    args = parser.parse_args()
    rows = json.loads((ROOT/"data/literature.json").read_text(encoding="utf-8"))
    groups = {
        "literature":rows,
        "historical_public":[r for r in json.loads((ROOT/"archive/imports/web_2026-10-15_data_literature.json").read_text(encoding="utf-8")) if r.get("provenance","").startswith("2026-10-14")],
        "new_handwriting":[r for r in rows if r["id"].startswith("N")],
        "umm":[r for r in rows if r["scope"]=="umm"],
        "resources":[r for r in rows if r["record_type"]=="resource"],
        "local_additions":[r for r in rows if r["id"].startswith("L")],
    }
    for name,recs in groups.items():
        (ROOT/f"data/{name}.json").write_text(json.dumps(recs,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        with (ROOT/f"data/{name}.csv").open("w",encoding="utf-8-sig",newline="") as f:
            w=csv.writer(f);w.writerow(HEADER);w.writerows(cells(r) for r in recs)
    targets = {
        "docs/all_literature.md":(rows,"全部公开文献与资源"),
        "docs/umm/literature.md":(groups["umm"],"UMM架构、协同、推理与效率文献表"),
        "docs/handwriting/new_literature.md":(groups["new_handwriting"],"书写过程重点文献"),
        "docs/handwriting/literature.md":([r for r in rows if r["scope"]!="umm"],"中文书写及相关研究"),
        "docs/local_additions.md":(groups["local_additions"],"补充文献与资源"),
    }
    for path,(recs,title) in targets.items():
        (ROOT/path).write_text(md(recs,title,path),encoding="utf-8")
    sheets = [("全部记录",rows),("UMM及机制",groups["umm"]),("书写及相关",[r for r in rows if r["scope"]!="umm"]),("资源",groups["resources"])]
    if args.xlsx_payload:
        args.xlsx_payload.write_text(json.dumps({title:[HEADER]+[cells(r) for r in recs] for title,recs in sheets},ensure_ascii=False),encoding="utf-8")
    if args.skip_xlsx:
        print(f"Rendered {len(rows)} catalog records; XLSX preserved for separate editing.")
        return
    workbook=Workbook()
    workbook.remove(workbook.active)
    for title,recs in sheets:
        sheet=workbook.create_sheet(title)
        sheet.append(HEADER)
        for r in recs:
            sheet.append(cells(r))
        sheet.freeze_panes="D2"
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="284B63")
        widths=[12,24,65,32,46,65,45,65,70,58,32,22,46,58,65,65,65,65,65,75,55,32]
        for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical="top",wrap_text=True)
    candidates=json.loads((ROOT/"data/discovery/arxiv_candidates.json").read_text(encoding="utf-8"))
    sheet=workbook.create_sheet("候选_仅发现")
    sheet.append(["arXiv","题名","版本","首次公开","更新","来源","状态"])
    for r in candidates:sheet.append([r['arxiv_id'],r['title'],r.get('version'),r.get('published'),r.get('updated'),r['url'],"检索元数据，不等于读过正文"])
    sheet.freeze_panes="C2";sheet.auto_filter.ref=sheet.dimensions
    sheet.column_dimensions['B'].width=90;sheet.column_dimensions['F'].width=55
    workbook.save(ROOT/"data/literature_public.xlsx")
    print(f"Rendered {len(rows)} catalog records and {len(candidates)} discovery records.")

if __name__ == "__main__":
    main()
